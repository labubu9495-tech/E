"""Place the verified working manuscript inside the user's Word template."""
from pathlib import Path
from copy import deepcopy
import json,re,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
SOURCE=OUT.parent/'paper_evidence_draft'
TEMPLATE=Path('/home/daji/users/dinglina/.codex/attachments/8baf34f0-bfe3-4793-bfc8-c55a4f924b7c/数学建模word模板.docx')
blocks=json.loads((SOURCE/'manuscript_blocks.json').read_text())
source=Document(SOURCE/'E题论文_论证增强工作稿.docx')
doc=Document(TEMPLATE)
front=[deepcopy(e) for e in list(doc._element.body)[:7]]
abstract_banner=[deepcopy(doc.paragraphs[i]._p) for i in [11,12,13]]
abstract_title=deepcopy(doc.paragraphs[17]._p)
sect=doc._element.body.sectPr
for e in list(doc._element.body):
 if e is not sect:doc._element.body.remove(e)
# Keep the actual cover graphics; remove stale AxMath/SEQ fields above them.
for el in list(front[0]):
 if el.tag==qn('w:pPr'):continue
 if el.tag==qn('w:r') and (el.xpath('.//w:drawing') or el.xpath('.//w:pict')):
  for child in list(el):
   if child.tag not in (qn('w:rPr'),qn('w:drawing'),qn('w:pict')):el.remove(child)
 else:front[0].remove(el)
for el in front:doc._element.body.insert(len(doc._element.body)-1,el)

def no_indent(p):
 p.paragraph_format.first_line_indent=Pt(0)
 pr=p._p.get_or_add_pPr();ind=pr.find(qn('w:ind'))
 if ind is not None:
  ind.set(qn('w:firstLineChars'),'0');ind.set(qn('w:leftChars'),'0')

def font(r,size=None,bold=None):
 if size is not None:r.font.size=Pt(size)
 if bold is not None:r.bold=bold
 r.font.name='Times New Roman';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')

def no_numbering(p):
 n=OxmlElement('w:numPr');i=OxmlElement('w:numId');i.set(qn('w:val'),'0');n.append(i);p._p.get_or_add_pPr().append(n)

def paragraph(text='',style='建模正文'):
 p=doc.add_paragraph(text,style=style)
 p.paragraph_format.widow_control=True
 return p

def clear_footer(s):
 s.footer.is_linked_to_previous=False
 for e in list(s.footer._element):s.footer._element.remove(e)
 s.footer.add_paragraph()

# Template defines A4, 4.226cm vertical margins, 2.207cm lateral margins.
# Use its styles and page dimensions; repair only stale numbering artifacts.
for name in ['Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name]
 pr=st.element.get_or_add_pPr()
 for n in list(pr):
  if n.tag==qn('w:numPr'):pr.remove(n)
 st.paragraph_format.keep_with_next=True
# Equation and TOC styles are derived from the template styles.
for name,base,size in [('工作稿公式','建模正文',11),('工作稿目录1','Normal',11),('工作稿目录2','Normal',10.5)]:
 if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st=doc.styles[name];st.base_style=doc.styles[base];st.font.name='Times New Roman';st.font.size=Pt(size)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
 st.paragraph_format.first_line_indent=Pt(0)
 st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(3)
 st.paragraph_format.line_spacing=1.0
 if '目录' in name:st.paragraph_format.keep_with_next=False
clear_footer(doc.sections[0])
# Abstract follows the template's second-page competition banner.
s=doc.add_section(WD_SECTION_START.NEW_PAGE);clear_footer(s)
for e in abstract_banner:doc._element.body.insert(len(doc._element.body)-1,e)
p=paragraph('题目：复杂场景下多模态情感预测的数学建模与算法设计');no_indent(p)
p.paragraph_format.space_before=Pt(14);p.paragraph_format.space_after=Pt(8)
for r in p.runs:font(r,14,True)
doc._element.body.insert(len(doc._element.body)-1,abstract_title)
abstract=[
 '针对多模态时间尺度差异、局部观测缺失与预测依据难以追溯的问题，本文以赛题附件为基础开展数据审查与建模。首先区分内容位置、原始可用观测和人工遮挡，修正专项文本缺失标记的处理，并统一训练与推理输入。第一问给出词级时间定位和区间重叠聚合方案，已核查100条原视频，自主特征与对齐实验尚待完成。',
 '第二问在3395条训练、728条验证样本上，比较池化融合、观测状态递推及MulT、EMT-DLFR适配模型，分析连续缺失与核心机制。按预定三种子四条件规则选中EMT-DLFR，其未追加人工遮挡的集成Macro-F1为0.5548、MAE为0.6348。相较MLP，MAE降低0.0240，但Accuracy略降；恢复约束与递推机制均未在所有条件稳定占优。文本50%遮挡时，分类与回归表现亦呈现不同取舍。',
 '现已完成附件3全部30条预测及复现检查。第三问提出模态子集扰动、局部删除和原素材定位的候选验证方案，尚无专项解释实测。本稿将数学依据、开发验证结论和待实施设计分别陈述，为后续完善三问提供可核验基础。'
]
for text in abstract:
 p=paragraph(text);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.1
 for r in p.runs:font(r,10.5)
p=paragraph('关键词：多模态情感预测；连续局部缺失；观测状态递推；跨模态恢复；证据追溯');no_indent(p)
for r in p.runs:font(r,10.5,True)
p=paragraph('阶段性工作稿：第一问与第三问尚待补齐实测成果。');no_indent(p)
for r in p.runs:font(r,9)
# Cached automatic TOC with bookmark page references. A second build fills
# page numbers from the rendered PDF, while Word can update the TOC later.
s=doc.add_section(WD_SECTION_START.NEW_PAGE);clear_footer(s)
p=doc.add_paragraph('目录',style='Heading 1');no_indent(p)
pagefile=OUT/'toc_pages.json';pages=json.loads(pagefile.read_text()) if pagefile.exists() else {}
toc=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');part=OxmlElement('w:docPartObj');gallery=OxmlElement('w:docPartGallery');gallery.set(qn('w:val'),'Table of Contents');part.append(gallery);pr.append(part);toc.append(pr);content=OxmlElement('w:sdtContent');toc.append(content)
headings=[x for x in blocks if x[0]=='h' and x[1]<=2]
bookmarks={x[2]:f'_PaperHeading{i:03d}' for i,x in enumerate(headings)}

def field_run(parent,kind,text=None):
 r=OxmlElement('w:r')
 if kind=='instr':
  t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=text;r.append(t)
 elif kind=='text':
  t=OxmlElement('w:t');t.text=text;r.append(t)
 else:
  t=OxmlElement('w:fldChar');t.set(qn('w:fldCharType'),kind);r.append(t)
 parent.append(r)
for i,x in enumerate(headings):
 p=doc.add_paragraph(style=f'工作稿目录{x[1]}');no_indent(p)
 p.paragraph_format.left_indent=Cm(.4 if x[1]==2 else 0)
 p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 if i==0:
  field_run(p._p,'begin');field_run(p._p,'instr',' TOC \\o "1-2" \\h \\z \\u ');field_run(p._p,'separate')
 p.add_run(x[2]);p.add_run('\t')
 field_run(p._p,'begin');field_run(p._p,'instr',f' PAGEREF {bookmarks[x[2]]} \\h ');field_run(p._p,'separate');field_run(p._p,'text',str(pages.get(x[2],'')));field_run(p._p,'end')
 if i==len(headings)-1:field_run(p._p,'end')
 content.append(p._p)
doc._element.body.insert(len(doc._element.body)-1,toc)
# Main text: restart visible page numbering at 1.
s=doc.add_section(WD_SECTION_START.NEW_PAGE);clear_footer(s);s.footer_distance=Cm(.8)
pg=s._sectPr.find(qn('w:pgNumType'))
if pg is None:pg=OxmlElement('w:pgNumType');s._sectPr.append(pg)
pg.set(qn('w:start'),'1')
p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;no_indent(p)
field_run(p._p,'begin');field_run(p._p,'instr',' PAGE ');field_run(p._p,'separate');field_run(p._p,'text','1');field_run(p._p,'end')
# Preserve exactly the checked editable equations from the source DOCX.
eqnodes={}
for p in source.paragraphs:
 if p._p.xpath('.//m:oMath'):
  key=re.search(r'（(\d+-\d+)）',p.text)
  assert key,p.text
  eqnodes[key[1]]=deepcopy(p._p.xpath('.//m:oMath')[0])

for x in blocks:
 kind=x[0]
 if kind=='h':
  p=doc.add_paragraph(x[2],style=f'Heading {x[1]}');no_indent(p)
  if x[1]==3:p.paragraph_format.left_indent=Cm(.25)
  if x[2] in bookmarks:
   i=list(bookmarks).index(x[2])+1
   a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(i));a.set(qn('w:name'),bookmarks[x[2]])
   z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(i));p._p.insert(1,a);p._p.append(z)
 elif kind in ('p','note'):
  p=paragraph(x[1])
  if kind=='note':
   for r in p.runs:r.bold=True
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');p._p.get_or_add_pPr().append(sh)
 elif kind=='eq':
  p=doc.add_paragraph(style='工作稿公式');no_indent(p);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
  p._p.append(deepcopy(eqnodes[x[2]]));font(p.add_run('  （'+x[2]+'）'),10)
 elif kind=='table':
  p=doc.add_paragraph(x[1],style='表名');no_indent(p);no_numbering(p);p.paragraph_format.keep_with_next=True
  for r in p.runs:font(r,10.5,True)
  t=doc.add_table(rows=1,cols=len(x[2]));t.autofit=True
  t.alignment=1
  # Three-line tables from the template: top/header/bottom, no vertical rules.
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'10');e.set(qn('w:color'),'000000');borders.append(e)
  t._tbl.tblPr.append(borders)
  for j,s in enumerate(x[2]):t.rows[0].cells[j].text=s
  for row in x[3]:
   cells=t.add_row().cells
   for j,s in enumerate(row):cells[j].text=s
  for i,row in enumerate(t.rows):
   rp=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');rp.append(ns)
   if i==0:rp.append(OxmlElement('w:tblHeader'))
   for cell in row.cells:
    tcpr=cell._tc.get_or_add_tcPr()
    if i==0:
     cb=OxmlElement('w:tcBorders');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'6');cb.append(bottom);tcpr.append(cb)
    for p in cell.paragraphs:
     p.style=doc.styles['建模正文'];no_indent(p);p.paragraph_format.left_indent=Pt(0)
     p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
     for r in p.runs:font(r,9.5 if len(x[2])>=5 else 10.5,i==0)
 elif kind=='fig':
  p=paragraph();no_indent(p);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  p.add_run().add_picture(x[1],width=Cm(min(x[3],15.8)))
  p=doc.add_paragraph(x[2],style='图名');no_indent(p);no_numbering(p)
  for r in p.runs:font(r,10.5,True)
 else:raise ValueError(kind)
# Keep the template author fields blank, and do not retain unused demo objects.
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.title='复杂场景下多模态情感预测的数学建模与算法设计（工作稿）'
doc.core_properties.subject='按用户数学建模Word模板排版的论证增强工作稿'
used=set(doc._element.xpath('//@r:embed'))|set(doc._element.xpath('//@r:id'))|set(doc._element.xpath('//@r:link'))
for rid,rel in list(doc.part.rels.items()):
 if rid not in used and (rel.reltype.endswith('/image') or rel.reltype.endswith('/oleObject')):doc.part.drop_rel(rid)
# Cached TOC is current; let the user update it after editing, not on every load.
u=doc.settings.element.find(qn('w:updateFields'))
if u is not None:doc.settings.element.remove(u)
path=OUT/'数学建模论文_套用模板工作稿.docx';doc.save(path)
print(json.dumps({'file':str(path),'sections':len(doc.sections),'tables':len(doc.tables),'inline_images':len(doc.inline_shapes),'equations':len(eqnodes),'toc_entries':len(headings),'filled_toc_pages':len(pages)},ensure_ascii=False))
